"""The text-box Layout: the style matrix's third column, observed through
``all_text``, python-docx's own view of the body, and the manifest."""

import io

import pytest
from docx import Document
from docx_text import all_text
from lxml import etree

from cvr.golden import CANDIDATES_DIR, LAYOUTS, Candidate, load_candidates

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
TXBX_CONTENT = f"{{{W}}}txbxContent"
W_T = f"{{{W}}}t"

CANDIDATES = load_candidates(CANDIDATES_DIR)


def c01() -> Candidate:
    (candidate,) = [c for c in CANDIDATES if c.id == "c01"]
    return candidate


def text_box():
    (layout,) = [layout for layout in LAYOUTS if layout.name == "text-box"]
    return layout


def body_texts(document: bytes) -> list[str]:
    """What python-docx can see: body paragraphs and table cells, no text boxes."""
    loaded = Document(io.BytesIO(document))
    texts = [paragraph.text for paragraph in loaded.paragraphs]
    for table in loaded.tables:
        for row in table.rows:
            for cell in row.cells:
                texts += [paragraph.text for paragraph in cell.paragraphs]
    return texts


def text_box_texts(document: bytes) -> list[list[str]]:
    """The ``w:t`` strings of each ``w:txbxContent`` in the main document part."""
    with io.BytesIO(document) as raw:
        loaded = Document(raw)
    return [
        [element.text or "" for element in box.iter(W_T)]
        for box in loaded.element.iter(TXBX_CONTENT)
    ]


def contact_strings(candidate: Candidate) -> list[str]:
    pii = candidate.pii
    values = [pii.phone, pii.email, *pii.address, *pii.urls, pii.dob]
    values += [pii.personal.nationality, pii.personal.marital_status]
    return [value for value in values if value]


# --- Text boxes


@pytest.mark.parametrize("candidate", CANDIDATES, ids=[c.id for c in CANDIDATES])
def test_contact_block_lives_in_a_text_box_not_the_body(candidate):
    document = text_box().generate(candidate).document
    body = body_texts(document)
    everywhere = all_text(document)
    for value in contact_strings(candidate):
        assert not any(value in text for text in body), value
        assert any(value in text for text in everywhere), value


def test_contact_block_and_skills_are_real_txbx_content_elements():
    candidate = c01()
    boxes = text_box_texts(text_box().generate(candidate).document)
    assert len(boxes) == 2
    contact, skills = boxes
    for value in contact_strings(candidate):
        assert any(value in text for text in contact), value
    for skill in candidate.content.skills:
        assert any(skill in text for text in skills), skill


def test_document_xml_is_well_formed_and_reloads():
    document = text_box().generate(c01()).document
    with io.BytesIO(document) as raw:
        loaded = Document(raw)
    etree.tostring(loaded.element)
    assert loaded.paragraphs


# --- Date style: ``Jan '20``


def c01_with(**first_experience_overrides) -> Candidate:
    data = c01().model_dump()
    data["content"]["experience"][0].update(first_experience_overrides)
    data["tags"] = []
    return Candidate.model_validate(data)


def test_c01_dates_print_as_abbreviated_month_and_two_digit_year():
    generated = text_box().generate(c01())
    (start,) = [
        d
        for d in generated.manifest.dates
        if d.section == "experience" and d.entry == 0 and d.which == "start"
    ]
    assert start.printed == "Mar '22"
    assert "Mar '22 - Jul '26" in all_text(generated.document)


def test_a_year_only_date_prints_as_the_full_year():
    candidate = c01_with(start={"year": 2020, "expected": "2020"})
    assert "2020 - Jul '26" in all_text(text_box().generate(candidate).document)


# --- Heading vocabulary, section order and bullets


def test_uses_its_heading_vocabulary_with_education_after_experience():
    texts = all_text(text_box().generate(c01()).document)
    headings = [
        "About Me",
        "Core Competencies",
        "Professional Experience",
        "Education & Training",
    ]
    positions = [texts.index(heading) for heading in headings]
    assert positions == sorted(positions)


def test_bullets_are_literal_hyphens_in_the_text():
    candidate = c01()
    texts = all_text(text_box().generate(candidate).document)
    first_bullet = candidate.content.experience[0].bullets[0]
    assert any(text.startswith("- ") and first_bullet in text for text in texts)
    assert not any(text.startswith("•") for text in texts)


# --- Scramble: experience rotated by one, education reversed


def titles_in_document_order(candidate: Candidate, document: bytes) -> list[int]:
    texts = all_text(document)
    titles = [entry.title for entry in candidate.content.experience]
    return sorted(range(len(titles)), key=lambda index: texts.index(titles[index]))


def test_c01_experience_is_rotated_by_one_and_education_reversed():
    candidate = c01()
    generated = text_box().generate(candidate)
    assert generated.manifest.experience_order == [1, 2, 0]
    assert generated.manifest.education_order == [1, 0]
    assert titles_in_document_order(candidate, generated.document) == [1, 2, 0]
    texts = all_text(generated.document)
    institutions = [entry.institution for entry in candidate.content.education]
    assert texts.index(institutions[1]) < texts.index(institutions[0])


def test_rotation_covers_only_dated_entries_and_undated_still_come_last():
    data = c01().model_dump()
    data["content"]["experience"][1].update(start=None, end=None)
    data["tags"] = []
    candidate = Candidate.model_validate(data)
    generated = text_box().generate(candidate)
    assert generated.manifest.experience_order == [2, 0, 1]
    assert titles_in_document_order(candidate, generated.document) == [2, 0, 1]
