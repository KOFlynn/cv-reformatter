"""The parser observed from the outside: a source document in, blocks out,
proven the way the Layouts were proven: against the dumb ``all_text`` helper
over every committed generated document, and against hand-made documents
for what the golden set does not exercise."""

import io
import json
import subprocess
import sys
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_SECTION
from docx_text import all_text, image_count

from cvr.golden import LAYOUTS as REGISTERED_LAYOUTS
from cvr.golden.generate import GENERATED_DIR
from cvr.golden.layouts.photo import placeholder_photo
from cvr.models import Image, NormalisationRule, Removal, RemovalRule, SourceBlock
from cvr.parse import ParsedDocument, parse
from cvr.text import CONFUSABLES, canonicalise

DOCUMENTS = sorted(GENERATED_DIR.glob("*.docx"))
STEMS = [path.stem for path in DOCUMENTS]
LAYOUTS = [layout.name for layout in REGISTERED_LAYOUTS]
CANDIDATE_IDS = sorted({stem.split("__")[0] for stem in STEMS})

# The confusable table's deletion rows: what NORM_INVISIBLE strips.
INVISIBLE = {char for char, replacement in CONFUSABLES.items() if replacement == ""}
CURLY_APOSTROPHE = "\u2019"
NBSP = "\u00a0"


def document(stem: str) -> Path:
    return GENERATED_DIR / f"{stem}.docx"


def manifest(stem: str) -> dict:
    return json.loads((GENERATED_DIR / f"{stem}.manifest.json").read_text("utf-8"))


def parsed(stem: str) -> ParsedDocument:
    return parse(document(stem).read_bytes())


def to_bytes(built: DocumentType) -> bytes:
    raw = io.BytesIO()
    built.save(raw)
    return raw.getvalue()


def test_the_golden_set_is_all_forty_eight_documents():
    assert len(DOCUMENTS) == 48
    assert CANDIDATE_IDS == [f"c{n:02d}" for n in range(1, 13)]


# --- Coverage: every text run is in some block, every block is in the text.


@pytest.mark.slow
@pytest.mark.parametrize("stem", STEMS)
def test_every_text_run_all_text_sees_is_in_some_block(stem):
    # A ``w:t`` never crosses a paragraph, so each one is inside one block.
    blocks = [canonicalise(block.text) for block in parsed(stem).blocks]
    missing = [
        text
        for text in map(canonicalise, all_text(document(stem)))
        if text and not any(text in block for block in blocks)
    ]
    assert not missing, f"{stem}: text the parser did not see: {missing}"


@pytest.mark.slow
@pytest.mark.parametrize("stem", STEMS)
def test_every_block_is_in_the_text_all_text_sees(stem):
    texts = " ".join(canonicalise(text) for text in all_text(document(stem)))
    for block in parsed(stem).blocks:
        assert isinstance(block, SourceBlock)
        assert block.text.strip(), f"{stem}: empty block {block.id}"
        assert canonicalise(block.text) in texts, f"{stem}: {block.id} not in text"


# --- Ids: addresses, stable for a file, different between Layouts.


@pytest.mark.slow
@pytest.mark.parametrize("stem", STEMS)
def test_two_parses_of_one_file_give_identical_blocks(stem):
    assert parsed(stem) == parsed(stem)


@pytest.mark.slow
@pytest.mark.parametrize("stem", STEMS)
def test_ids_are_unique_and_follow_the_grammar(stem):
    blocks = parsed(stem).blocks
    ids = [block.id for block in blocks]
    assert len(ids) == len(set(ids))
    for block in blocks:
        head, *rest = block.id.split(":")
        assert head == block.kind
        if head == "body":
            assert len(rest) == 1 and rest[0].isdigit()
        elif head == "table":
            n, r, c, p = rest
            assert n.isdigit() and p.isdigit()
            assert r[0] == "r" and r[1:].isdigit()
            assert c[0] == "c" and c[1:].isdigit()
        elif head == "textbox":
            n, b, p = rest
            assert n.isdigit() and b.isdigit() and p.isdigit()
        else:
            s, t, p = rest
            assert s.isdigit() and p.isdigit()
            assert t in ("default", "first", "even")


@pytest.mark.parametrize("candidate_id", CANDIDATE_IDS)
def test_ids_differ_between_the_four_layouts_of_one_candidate(candidate_id):
    # Addresses, not content hashes: the same Candidate laid out four ways
    # lives at four different sets of addresses.
    by_layout = [
        [block.id for block in parsed(f"{candidate_id}__{layout}").blocks]
        for layout in LAYOUTS
    ]
    for i, first in enumerate(by_layout):
        for second in by_layout[i + 1 :]:
            assert first != second


def test_gaps_are_real_an_empty_paragraph_keeps_its_index():
    built = Document()
    built.add_paragraph("First")
    built.add_paragraph()
    built.add_paragraph("Third")
    blocks = parse(to_bytes(built)).blocks
    assert [(b.id, b.text) for b in blocks] == [
        ("body:0", "First"),
        ("body:2", "Third"),
    ]


def test_a_table_contributes_its_cells_row_by_row_paragraphs_in_order():
    built = Document()
    built.add_paragraph("Before")
    table = built.add_table(rows=2, cols=2)
    table.cell(0, 0).paragraphs[0].text = "r0c0"
    table.cell(0, 1).paragraphs[0].text = "r0c1 first"
    table.cell(0, 1).add_paragraph("r0c1 second")
    table.cell(1, 0).paragraphs[0].text = "r1c0"
    built.add_paragraph("After")
    blocks = parse(to_bytes(built)).blocks
    assert [(b.id, b.text, b.kind) for b in blocks] == [
        ("body:0", "Before", "body"),
        ("table:1:r0:c0:0", "r0c0", "table"),
        ("table:1:r0:c1:0", "r0c1 first", "table"),
        ("table:1:r0:c1:1", "r0c1 second", "table"),
        ("table:1:r1:c0:0", "r1c0", "table"),
        ("body:2", "After", "body"),
    ]


# --- Text boxes: at their anchor's position, in anchor order.


@pytest.mark.parametrize("candidate_id", CANDIDATE_IDS)
def test_text_box_blocks_sit_at_their_anchor_position(candidate_id):
    blocks = parsed(f"{candidate_id}__text-box").blocks
    ids = [block.id for block in blocks]
    boxed = [block for block in blocks if block.kind == "textbox"]
    assert boxed, "the text-box Layout writes at least the contact box"
    # The name is the first body paragraph, the contact box is anchored in the
    # second (an otherwise empty paragraph, so no ``body:1``), then the rest.
    assert ids[0] == "body:0"
    assert ids[1].startswith("textbox:1:0:")
    assert "body:1" not in ids
    assert ids[len([i for i in ids if i.startswith("textbox:1:")]) + 1] == "body:2"
    # Every text box's blocks are contiguous and follow the last block of any
    # earlier body child: anchor order.
    anchors = [int(block.id.split(":")[1]) for block in boxed]
    assert anchors == sorted(anchors)
    for position, block in enumerate(blocks):
        if block.kind == "textbox":
            anchor = int(block.id.split(":")[1])
            assert all(
                int(earlier.id.split(":")[1]) <= anchor
                for earlier in blocks[:position]
                if earlier.kind in ("body", "table", "textbox")
            )


def test_text_box_paragraphs_are_numbered_within_their_box():
    blocks = parsed("c01__text-box").blocks
    contact = [block.id for block in blocks if block.id.startswith("textbox:1:0:")]
    assert contact == [f"textbox:1:0:{p}" for p in range(len(contact))]
    assert len(contact) > 1


@pytest.mark.parametrize("layout", ["single-column", "two-column", "header-footer"])
def test_other_layouts_have_no_text_box_blocks(layout):
    assert not [b for b in parsed(f"c01__{layout}").blocks if b.kind == "textbox"]


# --- Headers and footers: the type slot.


@pytest.mark.parametrize("candidate_id", CANDIDATE_IDS)
def test_header_footer_layout_carries_default_in_the_type_slot(candidate_id):
    blocks = parsed(f"{candidate_id}__header-footer").blocks
    headers = [b for b in blocks if b.kind == "header"]
    footers = [b for b in blocks if b.kind == "footer"]
    assert headers and footers
    for block in headers + footers:
        assert block.id.split(":")[1:3] == ["0", "default"]
    # Headers and footers come after the body, header before footer.
    kinds = [b.kind for b in blocks]
    assert kinds == sorted(kinds, key=["body", "header", "footer"].index)


@pytest.mark.parametrize("layout", ["single-column", "two-column", "text-box"])
def test_other_layouts_have_no_header_or_footer_blocks(layout):
    blocks = parsed(f"c01__{layout}").blocks
    assert not [b for b in blocks if b.kind in ("header", "footer")]


def test_a_first_page_header_yields_first_and_an_even_header_even():
    built = Document()
    built.add_paragraph("Body")
    section = built.sections[0]
    section.different_first_page_header_footer = True
    built.settings.odd_and_even_pages_header_footer = True
    section.header.paragraphs[0].text = "Default header"
    section.first_page_header.paragraphs[0].text = "First header"
    section.even_page_header.paragraphs[0].text = "Even header"
    section.footer.paragraphs[0].text = "Default footer"
    section.first_page_footer.paragraphs[0].text = "First footer"
    blocks = parse(to_bytes(built)).blocks
    assert [(b.id, b.text) for b in blocks] == [
        ("body:0", "Body"),
        ("header:0:default:0", "Default header"),
        ("header:0:first:0", "First header"),
        ("header:0:even:0", "Even header"),
        ("footer:0:default:0", "Default footer"),
        ("footer:0:first:0", "First footer"),
    ]


def test_a_second_section_numbers_its_own_header():
    built = Document()
    built.add_paragraph("One")
    built.sections[0].header.paragraphs[0].text = "Header one"
    second = built.add_section(WD_SECTION.NEW_PAGE)
    built.add_paragraph("Two")
    second.header.is_linked_to_previous = False
    second.header.paragraphs[0].text = "Header two"
    blocks = parse(to_bytes(built)).blocks
    assert [(b.id, b.text) for b in blocks if b.kind == "header"] == [
        ("header:0:default:0", "Header one"),
        ("header:1:default:0", "Header two"),
    ]


def test_a_linked_header_is_not_read_twice():
    built = Document()
    built.add_paragraph("One")
    built.sections[0].header.paragraphs[0].text = "Shared header"
    built.add_section(WD_SECTION.NEW_PAGE)
    built.add_paragraph("Two")
    blocks = parse(to_bytes(built)).blocks
    assert [b.id for b in blocks if b.kind == "header"] == ["header:0:default:0"]


# --- NORM_INVISIBLE: the deletion rows stripped and logged, nothing else touched.


@pytest.mark.slow
@pytest.mark.parametrize("stem", STEMS)
def test_invisible_characters_are_stripped_and_every_one_is_logged(stem):
    result = parsed(stem)
    joined = "".join(all_text(document(stem)))
    for block in result.blocks:
        assert not INVISIBLE & set(block.text), f"{stem}: {block.id} keeps invisibles"
    logged = [char for event in result.normalisations for char in event.characters]
    assert sorted(logged) == sorted(char for char in joined if char in INVISIBLE)
    ids = {block.id for block in result.blocks}
    for event in result.normalisations:
        assert event.rule is NormalisationRule.INVISIBLE
        assert event.block_id in ids  # the golden set has no invisible-only line
        assert event.characters, "an event with nothing stripped"
    assert len({event.block_id for event in result.normalisations}) == len(
        result.normalisations
    ), "one event per block"


@pytest.mark.slow
@pytest.mark.parametrize("stem", STEMS)
def test_events_match_the_manifest_deletion_row_confusables(stem):
    injected = {
        chr(int(code.removeprefix("U+"), 16)) for code in manifest(stem)["confusables"]
    } & INVISIBLE
    stripped = {c for event in parsed(stem).normalisations for c in event.characters}
    assert stripped == injected
    if stem.endswith("__text-box"):
        assert stripped == {"\u00ad"}
    if stem.endswith("__header-footer"):
        assert stripped == {"\u200b"}
    if stem.endswith(("__single-column", "__two-column")):
        assert stripped == set()


@pytest.mark.parametrize("layout", ["two-column", "header-footer"])
def test_visible_confusables_survive(layout):
    stem = f"c01__{layout}"
    joined = "".join(all_text(document(stem)))
    blocks = "".join(block.text for block in parsed(stem).blocks)
    assert CURLY_APOSTROPHE in joined
    assert blocks.count(CURLY_APOSTROPHE) == joined.count(CURLY_APOSTROPHE)
    assert "'" not in blocks
    # Every substitution row the Layout injected is still there, uncounted by
    # the coverage tests because they canonicalise both sides.
    for code in manifest(stem)["confusables"]:
        char = chr(int(code.removeprefix("U+"), 16))
        if char not in INVISIBLE:
            assert blocks.count(char) == joined.count(char) > 0, code


def test_non_breaking_spaces_are_a_substitution_row_and_survive():
    stem = "c01__text-box"
    joined = "".join(all_text(document(stem)))
    blocks = "".join(block.text for block in parsed(stem).blocks)
    assert NBSP in joined
    assert blocks.count(NBSP) == joined.count(NBSP)


def test_a_paragraph_that_is_only_invisibles_is_a_gap_and_is_still_logged():
    # No block: there is nothing to label. Still an event: the characters were
    # taken out of the file, and the log is the only place that says so.
    built = Document()
    built.add_paragraph("\u200b\u00ad")
    built.add_paragraph("Kept\u200b")
    result = parse(to_bytes(built))
    assert [(b.id, b.text) for b in result.blocks] == [("body:1", "Kept")]
    assert [(e.block_id, e.characters) for e in result.normalisations] == [
        ("body:0", ["\u200b", "\u00ad"]),
        ("body:1", ["\u200b"]),
    ]


# --- Images: collected by content hash, removed under RM_PHOTO.


@pytest.mark.slow
@pytest.mark.parametrize("stem", STEMS)
def test_only_two_column_documents_carry_an_image_and_it_is_removed(stem):
    result = parsed(stem)
    expected = 1 if stem.endswith("__two-column") else 0
    assert len(result.images) == expected == image_count(document(stem))
    assert len(result.removals) == expected
    for image, removal in zip(result.images, result.removals, strict=True):
        assert isinstance(image, Image)
        assert image.part.startswith("word/media/")
        assert len(image.sha256) == 64 and image.size > 0
        assert removal == Removal(rule=RemovalRule.PHOTO, subject=image)


def test_the_placeholder_photo_hashes_the_same_in_every_two_column_document():
    hashes = {
        image.sha256
        for candidate_id in CANDIDATE_IDS
        for image in parsed(f"{candidate_id}__two-column").images
    }
    assert len(hashes) == 1


def test_an_image_in_a_header_is_collected_too():
    built = Document()
    built.add_paragraph("Body")
    header = built.sections[0].header
    header.paragraphs[0].add_run().add_picture(io.BytesIO(placeholder_photo()))
    result = parse(to_bytes(built))
    assert len(result.images) == 1
    assert [b.id for b in result.blocks] == ["body:0"]


def test_the_parse_package_depends_on_text_and_models_only():
    code = (
        "import sys, cvr.parse; "
        "sys.exit(any(m in sys.modules for m in ('cvr.golden', 'cvr.eval')))"
    )
    assert subprocess.run([sys.executable, "-c", code], check=False).returncode == 0
