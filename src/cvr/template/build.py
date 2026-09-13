"""Builds the Fictitious Recruitment template with python-docx.

NEVER HAND-EDIT THE OUTPUT (``templates/fictitious_recruitment.docx``). Word
splits ``{{ }}`` and ``{%p %}`` tags across runs as you type them, which
breaks docxtpl invisibly (ADR-0006). Change this script and regenerate with
``python -m cvr.template.build``; the test suite fails while the committed
file and a fresh build disagree.

Section order per the brief: header (wordmark and name), profile, key skills,
education, experience, certifications, additional information, footer, review
appendix. Every section is wrapped in a paragraph-level ``{%p if %}`` so an
empty one leaves no heading and no blank paragraph behind. Each tag sits in
exactly one run because python-docx writes a run per ``add_run`` call.

Branding is deliberately modest: a text wordmark, one accent colour, one
font, no images. The red review-appendix banner is the one visual element
that matters for the demo.
"""

import argparse
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Length, Pt, RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from cvr.template.paths import TEMPLATE_PATH

__all__ = ["BANNER", "FOOTER", "WORDMARK", "build_template", "main"]

WORDMARK = "Fictitious Recruitment"
BANNER = "TEXT NOT PLACED — NEEDS HUMAN REVIEW"
FOOTER = "References available on request"

FONT = "Calibri"
ACCENT = RGBColor(0x1F, 0x4E, 0x79)
RED = RGBColor(0xC0, 0x00, 0x00)


def build_template() -> DocumentType:
    """The template as a python-docx Document, tags and all."""
    document = Document()
    _style(document)
    section = document.sections[0]

    wordmark = section.header.paragraphs[0]
    wordmark.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(wordmark, WORDMARK, bold=True, colour=ACCENT, size=Pt(14))
    _run(section.footer.paragraphs[0], FOOTER, italic=True)
    section.footer.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    _content(document, "{{ name }}", style="Title")

    _list_section(document, "Profile", "profile")
    _list_section(document, "Key Skills", "skills", style="List Bullet")

    _tag(document, "{%p if education %}")
    document.add_heading("Education", level=1)
    _tag(document, "{%p for entry in education %}")
    _content(document, "{{ entry.qualification }}", bold=True)
    _content(document, "{{ entry.institution }}")
    _date_line(document)
    _tag(document, "{%p for line in entry.details %}")
    _content(document, "{{ line }}")
    _tag(document, "{%p endfor %}")
    _tag(document, "{%p endfor %}")
    _tag(document, "{%p endif %}")

    _tag(document, "{%p if experience %}")
    document.add_heading("Experience", level=1)
    _tag(document, "{%p for entry in experience %}")
    _content(document, "{{ entry.title }}", bold=True)
    _content(
        document,
        "{{ entry.employer }}{% if entry.location %}, {{ entry.location }}{% endif %}",
    )
    _date_line(document)
    _tag(document, "{%p for bullet in entry.bullets %}")
    _content(document, "{{ bullet }}", style="List Bullet")
    _tag(document, "{%p endfor %}")
    _tag(document, "{%p endfor %}")
    _tag(document, "{%p endif %}")

    _list_section(document, "Certifications", "certifications", style="List Bullet")
    _list_section(document, "Additional Information", "additional")

    # The review appendix: present only when something was not placed, and
    # unmissable when it is (ADR-0004, review by exception).
    _tag(document, "{%p if unplaced %}")
    _run(document.add_paragraph(), BANNER, bold=True, colour=RED, size=Pt(20))
    _tag(document, "{%p for fragment in unplaced %}")
    _content(document, "{{ fragment }}")
    _tag(document, "{%p endfor %}")
    _tag(document, "{%p endif %}")

    return document


def _style(document: DocumentType) -> None:
    # One font everywhere, one accent colour on the headings.
    for name in ("Normal", "Title", "Heading 1", "List Bullet"):
        document.styles[name].font.name = FONT
    document.styles["Normal"].font.size = Pt(11)
    for name in ("Title", "Heading 1"):
        document.styles[name].font.color.rgb = ACCENT


def _run(
    paragraph: Paragraph,
    text: str,
    *,
    bold: bool = False,
    italic: bool = False,
    colour: RGBColor | None = None,
    size: Length | None = None,
) -> Run:
    # Only set what is asked for: ``run.bold = False`` writes an explicit
    # off-toggle into the XML rather than nothing.
    run = paragraph.add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if colour is not None:
        run.font.color.rgb = colour
    if size is not None:
        run.font.size = size
    return run


def _tag(document: DocumentType, tag: str) -> None:
    """A paragraph-level docxtpl tag, alone in its paragraph as docxtpl requires."""
    document.add_paragraph().add_run(tag)


def _content(
    document: DocumentType, text: str, *, style: str | None = None, bold: bool = False
) -> None:
    """A content paragraph whose text is (or contains) inline Jinja expressions."""
    _run(document.add_paragraph(style=style), text, bold=bold)


def _list_section(
    document: DocumentType, heading: str, field: str, *, style: str | None = None
) -> None:
    """A heading over one paragraph per item of a list field, all of it
    absent when the list is empty."""
    _tag(document, f"{{%p if {field} %}}")
    document.add_heading(heading, level=1)
    _tag(document, f"{{%p for line in {field} %}}")
    _content(document, "{{ line }}", style=style)
    _tag(document, "{%p endfor %}")
    _tag(document, "{%p endif %}")


def _date_line(document: DocumentType) -> None:
    # ``MM/YYYY – MM/YYYY``; a present role's end reads ``Present`` because that
    # is its DateValue.expected. Omitted when neither date exists; one date
    # alone prints without the dash, mirroring the Layouts.
    _tag(document, "{%p if entry.start or entry.end %}")
    _content(
        document,
        "{% if entry.start %}{{ entry.start.expected }}{% endif %}"
        "{% if entry.start and entry.end %} – {% endif %}"
        "{% if entry.end %}{{ entry.end.expected }}{% endif %}",
    )
    _tag(document, "{%p endif %}")


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path, default=TEMPLATE_PATH)
    args = parser.parse_args(argv)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    build_template().save(args.out)
    print(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
