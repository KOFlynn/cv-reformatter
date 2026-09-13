"""The header/footer Layout: the style matrix's fourth column.

The contact block is not in the body at all: phone and email sit in the page
header, address and URLs in the page footer, so a pipeline that reads only
body paragraphs never sees them and leaks them. ``2020-01`` dates, literal
``–`` bullets, education at the very bottom, experience and education both
reversed (oldest first), zero-width spaces and curly apostrophes from the
shared confusable table.
"""

from typing import ClassVar

from docx import Document
from docx.document import Document as DocumentType

from cvr.golden.candidate import Candidate
from cvr.golden.layouts.base import (
    Decisions,
    FragmentPlacement,
    Layout,
    Plan,
    PlannedEntry,
    Section,
)
from cvr.models import EducationEntry, ExperienceEntry

__all__ = ["HeaderFooterLayout"]


class HeaderFooterLayout(Layout):
    name: ClassVar[str] = "header-footer"
    contact_block: ClassVar[str] = "header-footer"

    def format_date(self, month: int | None, year: int) -> str:
        return str(year) if month is None else f"{year}-{month:02d}"

    def scramble(self, section: Section, indices: list[int]) -> list[int]:
        # Both sections oldest first.
        return indices[::-1]

    def render(
        self, candidate: Candidate, plan: Plan, decisions: Decisions
    ) -> DocumentType:
        content, pii = candidate.content, candidate.pii
        document = Document()
        section = document.sections[0]

        # Header: phone and email. Footer: address and URLs. A fresh header or
        # footer already holds one empty paragraph; the first line goes into it.
        _lines(section.header, [pii.phone, pii.email])
        _lines(section.footer, [*pii.address, *pii.urls])

        body = document
        body.add_paragraph(content.name, style="Title")
        _lines(body, [pii.dob, pii.personal.nationality, pii.personal.marital_status])

        if content.profile:
            body.add_heading("Personal Statement", level=1)
            _lines(body, content.profile)

        if content.skills:
            body.add_heading("Technical Skills", level=1)
            _bullets(body, content.skills)

        if plan.experience:
            body.add_heading("Employment", level=1)
            for planned in plan.experience:
                _experience(body, planned)

        if content.certifications:
            body.add_heading("Certifications", level=1)
            _bullets(body, content.certifications)

        if content.additional:
            body.add_heading("Further Information", level=1)
            _lines(body, content.additional)

        if pii.referees:
            body.add_heading("Referees", level=1)
            for referee in pii.referees:
                _line(body, referee.name, bold=True)
                _lines(body, [referee.role, *referee.contact])

        # Education at the very bottom of the body, after everything else.
        if plan.education:
            body.add_heading("Qualifications", level=1)
            for planned in plan.education:
                _education(body, planned)

        # Unplaceable fragments trail the footer, under the contact lines: text
        # a pipeline reading only the body never sees, in a part it must still
        # sweep for the appendix.
        for index, fragment in enumerate(candidate.unplaceable):
            _line(section.footer, fragment)
            decisions.fragments.append(
                FragmentPlacement(index=index, location="footer-end")
            )

        return document


def _paragraph(container):
    """The next paragraph to write into: a header or footer's one empty
    paragraph while it is still empty, otherwise a fresh one."""
    paragraphs = container.paragraphs
    if paragraphs and not paragraphs[0].text and len(paragraphs) == 1:
        return paragraphs[0]
    return container.add_paragraph()


def _line(container, text: str, *, bold: bool = False, italic: bool = False) -> None:
    # Only set what is asked for: ``run.bold = False`` would write an explicit
    # off-toggle into the XML rather than nothing.
    run = _paragraph(container).add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True


def _lines(container, texts: list[str | None]) -> None:
    """One paragraph per text, skipping absent values."""
    for text in texts:
        if text:
            _line(container, text)


def _bullets(container, items: list[str]) -> None:
    # A literal en dash as the bullet glyph, in the text itself.
    for item in items:
        _line(container, f"– {item}")


def _date_line(container, planned: PlannedEntry) -> None:
    # ``2022-03 to 2026-07``: joined by a word, since a dash between
    # dash-separated dates would be ambiguous to read and the column's
    # confusables are the zero-width space and the curly apostrophe only.
    parts = [part for part in (planned.start, planned.end) if part is not None]
    if parts:
        _line(container, " to ".join(parts))


def _experience(container, planned: PlannedEntry[ExperienceEntry]) -> None:
    entry = planned.entry
    _line(container, entry.title, bold=True)
    employer_line = (
        entry.employer
        if entry.location is None
        else f"{entry.employer}, {entry.location}"
    )
    _line(container, employer_line, italic=True)
    _date_line(container, planned)
    _bullets(container, entry.bullets)


def _education(container, planned: PlannedEntry[EducationEntry]) -> None:
    entry = planned.entry
    _line(container, entry.qualification, bold=True)
    _line(container, entry.institution, italic=True)
    _date_line(container, planned)
    _bullets(container, entry.details)
