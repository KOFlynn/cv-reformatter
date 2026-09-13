"""The two-column table Layout: the style matrix's second column.

One two-cell table holds the whole CV. The left cell carries the contact
block, a generated placeholder photo and education; the right cell carries
the name, summary, skills and work history. ``01/2020`` dates, literal ``•``
bullets, experience reversed (oldest first), curly quotes and en dashes from
the shared confusable table. Reading order is a trap in itself: the pipeline
has to realise the columns are not one stream of text.
"""

import io
from typing import ClassVar

from docx import Document
from docx.document import Document as DocumentType
from docx.shared import Cm
from docx.table import _Cell

from cvr.golden.candidate import Candidate
from cvr.golden.layouts.base import (
    Decisions,
    FragmentPlacement,
    Layout,
    Plan,
    PlannedEntry,
    Section,
)
from cvr.golden.layouts.photo import placeholder_photo
from cvr.models import EducationEntry, ExperienceEntry

__all__ = ["TwoColumnLayout"]

_LEFT_WIDTH = Cm(5.5)
_RIGHT_WIDTH = Cm(11.5)
_PHOTO_WIDTH = Cm(3.2)
_EN_DASH = "–"
_APOSTROPHE = "’"
_OPEN_DOUBLE, _CLOSE_DOUBLE = "“", "”"


class TwoColumnLayout(Layout):
    name: ClassVar[str] = "two-column"
    contact_block: ClassVar[str] = "left-column"

    def format_date(self, month: int | None, year: int) -> str:
        return str(year) if month is None else f"{month:02d}/{year}"

    def scramble(self, section: Section, indices: list[int]) -> list[int]:
        # Experience oldest first; education stays in Candidate order.
        return indices[::-1] if section == "experience" else indices

    def render(
        self, candidate: Candidate, plan: Plan, decisions: Decisions
    ) -> DocumentType:
        content, pii = candidate.content, candidate.pii
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        left, right = table.rows[0].cells
        left.width, right.width = _LEFT_WIDTH, _RIGHT_WIDTH
        # Left column: photo, contact block, then education.
        left.add_paragraph().add_run().add_picture(
            io.BytesIO(placeholder_photo()), width=_PHOTO_WIDTH
        )
        decisions.photo = True

        _heading(left, "Contact")
        contact = [pii.phone, pii.email, *pii.address, *pii.urls, pii.dob]
        _lines(left, contact, decisions)
        _lines(left, [pii.personal.nationality, pii.personal.marital_status], decisions)

        if plan.education:
            _heading(left, "Academic Background")
            for planned in plan.education:
                _education(left, planned, decisions)

        if pii.referees:
            _heading(left, "Referees")
            for referee in pii.referees:
                _line(left, referee.name, decisions, bold=True)
                _lines(left, [referee.role, *referee.contact], decisions)

        # Right column: name, then the rest.
        right.add_paragraph(_curl(content.name, decisions), style="Title")

        if content.profile:
            _heading(right, "Summary")
            _lines(right, content.profile, decisions)

        if content.skills:
            _heading(right, "Skills")
            _bullets(right, content.skills, decisions)

        if plan.experience:
            _heading(right, "Work History")
            for planned in plan.experience:
                _experience(right, planned, decisions)

        if content.certifications:
            _heading(right, "Certifications")
            _bullets(right, content.certifications, decisions)

        if content.additional:
            _heading(right, "Other Information")
            _lines(right, content.additional, decisions)

        # Unplaceable fragments trail the left column, under the referees.
        for index, fragment in enumerate(candidate.unplaceable):
            _line(left, fragment, decisions)
            decisions.fragments.append(
                FragmentPlacement(index=index, location="left-column-end")
            )

        return document


def _heading(cell: _Cell, text: str) -> None:
    cell.add_paragraph(text, style="Heading 1")


def _curl(text: str, decisions: Decisions) -> str:
    """Word's smart quotes: every apostrophe becomes a right single quote and
    straight double quotes alternate open and close. Each curly character comes
    from the shared confusable table and is recorded as injected."""
    if "'" in text:
        decisions.injected(_APOSTROPHE)
        text = text.replace("'", _APOSTROPHE)
    if '"' in text:
        halves = text.split('"')
        text = halves[0]
        for position, half in enumerate(halves[1:]):
            quote = _OPEN_DOUBLE if position % 2 == 0 else _CLOSE_DOUBLE
            decisions.injected(quote)
            text += quote + half
    return text


def _line(
    cell: _Cell,
    text: str,
    decisions: Decisions | None = None,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    """One paragraph. Pass ``decisions`` to curl the text's quotes; date lines
    pass nothing so the document says what the manifest says was printed."""
    if decisions is not None:
        text = _curl(text, decisions)
    # Only set what is asked for: ``run.bold = False`` would write an explicit
    # off-toggle into the XML rather than nothing.
    run = cell.add_paragraph().add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True


def _lines(cell: _Cell, texts: list[str | None], decisions: Decisions) -> None:
    """One paragraph per text, skipping absent values."""
    for text in texts:
        if text:
            _line(cell, text, decisions)


def _bullets(cell: _Cell, items: list[str], decisions: Decisions) -> None:
    # A literal bullet glyph in the text itself, not Word list numbering.
    for item in items:
        _line(cell, f"• {item}", decisions)


def _date_line(cell: _Cell, planned: PlannedEntry, decisions: Decisions) -> None:
    """``start – end`` with an en dash from the shared table; a lone date has
    no dash and injects nothing."""
    parts = [part for part in (planned.start, planned.end) if part is not None]
    if len(parts) == 2:
        decisions.injected(_EN_DASH)
    if parts:
        _line(cell, f" {_EN_DASH} ".join(parts))


def _experience(
    cell: _Cell, planned: PlannedEntry[ExperienceEntry], decisions: Decisions
) -> None:
    entry = planned.entry
    _line(cell, entry.title, decisions, bold=True)
    employer_line = (
        entry.employer
        if entry.location is None
        else f"{entry.employer}, {entry.location}"
    )
    _line(cell, employer_line, decisions, italic=True)
    _date_line(cell, planned, decisions)
    _bullets(cell, entry.bullets, decisions)


def _education(
    cell: _Cell, planned: PlannedEntry[EducationEntry], decisions: Decisions
) -> None:
    entry = planned.entry
    _line(cell, entry.qualification, decisions, bold=True)
    _line(cell, entry.institution, decisions, italic=True)
    _date_line(cell, planned, decisions)
    _bullets(cell, entry.details, decisions)
