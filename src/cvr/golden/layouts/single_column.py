"""The single-column Layout: the clean control column of the style matrix.

``January 2020`` dates, Word list numbering for bullets, the contact block at
the top of the body, Candidate order throughout, no confusables, no photo.
Education sits after profile and skills. Anything the pipeline gets wrong
here it gets wrong on a tidy document, which is what a control is for.
"""

from typing import ClassVar

from docx import Document
from docx.document import Document as DocumentType

from cvr.golden.candidate import Candidate
from cvr.golden.layouts.base import (
    Decisions,
    FragmentPlacement,
    Layout,
    Plan,
    PlannedEntry,
)
from cvr.models import EducationEntry, ExperienceEntry

__all__ = ["SingleColumnLayout"]

_MONTHS = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]


class SingleColumnLayout(Layout):
    name: ClassVar[str] = "single-column"
    contact_block: ClassVar[str] = "body-top"

    def format_date(self, month: int | None, year: int) -> str:
        return str(year) if month is None else f"{_MONTHS[month - 1]} {year}"

    def render(
        self, candidate: Candidate, plan: Plan, decisions: Decisions
    ) -> DocumentType:
        content, pii = candidate.content, candidate.pii
        document = Document()
        document.add_paragraph(content.name, style="Title")

        # Contact block at the top of the body, one line per detail.
        _lines(document, [pii.phone, pii.email, *pii.address, *pii.urls, pii.dob])
        _lines(document, [pii.personal.nationality, pii.personal.marital_status])

        if content.profile:
            document.add_heading(decisions.heading("Profile"), level=1)
            _lines(document, content.profile)

        if content.skills:
            document.add_heading(decisions.heading("Key Skills"), level=1)
            _bullets(document, content.skills)

        if plan.education:
            document.add_heading(decisions.heading("Education"), level=1)
            for planned in plan.education:
                _education(document, planned)

        if plan.experience:
            document.add_heading(decisions.heading("Experience"), level=1)
            for planned in plan.experience:
                _experience(document, planned)

        if content.certifications:
            document.add_heading(decisions.heading("Certifications"), level=1)
            _bullets(document, content.certifications)

        if content.additional:
            document.add_heading(decisions.heading("Additional Information"), level=1)
            _lines(document, content.additional)

        if pii.referees:
            document.add_heading(decisions.heading("References"), level=1)
            for referee in pii.referees:
                _line(document, referee.name, bold=True)
                _lines(document, [referee.role, *referee.contact])

        # Unplaceable fragments trail the body as plain paragraphs.
        for index, fragment in enumerate(candidate.unplaceable):
            _line(document, fragment)
            decisions.fragments.append(
                FragmentPlacement(index=index, location="body-end")
            )

        return document


def _line(
    document: DocumentType, text: str, *, bold: bool = False, italic: bool = False
) -> None:
    # Only set what is asked for: ``run.bold = False`` would write an explicit
    # off-toggle into the XML rather than nothing.
    run = document.add_paragraph().add_run(text)
    if bold:
        run.bold = True
    if italic:
        run.italic = True


def _lines(document: DocumentType, texts: list[str | None]) -> None:
    """One paragraph per text, skipping absent values."""
    for text in texts:
        if text:
            _line(document, text)


def _bullets(document: DocumentType, items: list[str]) -> None:
    # Word list numbering: the glyph comes from the style, not the text.
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def _date_line(document: DocumentType, planned: PlannedEntry) -> None:
    parts = [part for part in (planned.start, planned.end) if part is not None]
    if parts:
        _line(document, " - ".join(parts))


def _experience(document: DocumentType, planned: PlannedEntry[ExperienceEntry]) -> None:
    entry = planned.entry
    _line(document, entry.title, bold=True)
    employer_line = (
        entry.employer
        if entry.location is None
        else f"{entry.employer}, {entry.location}"
    )
    _line(document, employer_line, italic=True)
    _date_line(document, planned)
    _bullets(document, entry.bullets)


def _education(document: DocumentType, planned: PlannedEntry[EducationEntry]) -> None:
    entry = planned.entry
    _line(document, entry.qualification, bold=True)
    _line(document, entry.institution, italic=True)
    _date_line(document, planned)
    _bullets(document, entry.details)
